#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
BA Markdown to Professional DOCX Exporter
Compliant with Vietnam Administrative Document Formatting Standards (Nghị định 30/2020/NĐ-CP)
Author: Senior AI Coding Agent
Standard: Senior Engineering Playbook (Build the Lever & Clean I/O)
"""

import sys
import os
import re
import argparse
from pathlib import Path

# Ensure docx is installed
try:
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsdecls, qn
except ImportError:
    print("[ERROR] python-docx is not installed. Please run: pip install python-docx")
    sys.exit(1)

# Configure UTF-8 encoding for Windows stdout
sys.stdout.reconfigure(encoding='utf-8')

def set_cell_background(cell, color_hex):
    """Sets cell shading background color."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets inner margins for a table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def apply_vietnam_nd30_formatting(doc):
    """
    Applies Vietnamese Administrative Formatting Standards (Nghị định 30/2020/NĐ-CP):
    - Paper: A4 (8.27 x 11.69 inches)
    - Margins: Top 20mm (0.79 in), Bottom 20mm (0.79 in), Left 30mm (1.18 in), Right 20mm (0.79 in)
    - Default Font: Times New Roman, 13pt
    - Line spacing: 1.3 lines (~16.9pt)
    """
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.79)     # 20 mm
        section.bottom_margin = Inches(0.79)  # 20 mm
        section.left_margin = Inches(1.18)    # 30 mm
        section.right_margin = Inches(0.79)   # 20 mm

    # Configure Normal Style
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(13)
    style_normal.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    style_normal.paragraph_format.line_spacing = 1.3
    style_normal.paragraph_format.space_after = Pt(4)

def parse_markdown_table(lines):
    """Parses markdown table lines into list of rows with clean text."""
    rows = []
    for line in lines:
        line = line.strip()
        if not line.startswith('|') or not line.endswith('|'):
            continue
        cells = [c.strip() for c in line.split('|')[1:-1]]
        # Skip separator rows like |---|---|
        if cells and all(re.match(r'^:?-+:?$', c) for c in cells):
            continue
        rows.append(cells)
    return rows

def convert_md_to_docx(md_path, docx_path):
    """Converts a Markdown file into a formatted Microsoft Word DOCX file."""
    md_file = Path(md_path)
    if not md_file.exists():
        print(f"[FAIL] Markdown file not found: {md_path}")
        return False

    raw_text = md_file.read_text(encoding='utf-8')
    lines = raw_text.splitlines()

    doc = docx.Document()
    apply_vietnam_nd30_formatting(doc)

    in_table = False
    table_lines = []
    in_code_block = False
    code_block_lines = []

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Handle Fenced Code Blocks (e.g. Mermaid or Gherkin)
        if stripped.startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_block_lang = stripped[3:].strip()
                code_block_lines = []
            else:
                in_code_block = False
                # Insert code block box
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                set_cell_background(cell, "F4F6F9")
                set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
                
                header_p = cell.paragraphs[0]
                header_p.text = f"[{code_block_lang.upper() if code_block_lang else 'CODE'}]"
                header_p.runs[0].font.name = 'Consolas'
                header_p.runs[0].font.size = Pt(9.5)
                header_p.runs[0].font.bold = True
                header_p.runs[0].font.color.rgb = RGBColor(0x02, 0x55, 0x88)
                header_p.paragraph_format.space_after = Pt(2)

                code_p = cell.add_paragraph()
                code_p.text = "\n".join(code_block_lines)
                code_p.runs[0].font.name = 'Consolas'
                code_p.runs[0].font.size = Pt(10)
                code_p.runs[0].font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                code_p.paragraph_format.line_spacing = 1.15
                code_p.paragraph_format.space_after = Pt(0)
                
                doc.add_paragraph() # Spacer
            i += 1
            continue

        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue

        # Handle Markdown Tables
        if stripped.startswith("|") and stripped.endswith("|"):
            in_table = True
            table_lines.append(stripped)
            i += 1
            continue
        elif in_table:
            # End of table detected
            in_table = False
            parsed_rows = parse_markdown_table(table_lines)
            table_lines = []
            if parsed_rows:
                num_rows = len(parsed_rows)
                num_cols = max(len(r) for r in parsed_rows)
                tbl = doc.add_table(rows=num_rows, cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                tbl.autofit = True

                for r_idx, row_data in enumerate(parsed_rows):
                    for c_idx in range(num_cols):
                        val = row_data[c_idx] if c_idx < len(row_data) else ""
                        cell = tbl.cell(r_idx, c_idx)
                        cell.text = val
                        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)

                        if r_idx == 0:
                            # Header Row
                            set_cell_background(cell, "003366") # Navy Blue
                            p = cell.paragraphs[0]
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            if p.runs:
                                p.runs[0].font.bold = True
                                p.runs[0].font.name = 'Times New Roman'
                                p.runs[0].font.size = Pt(11.5)
                                p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        else:
                            # Data Rows
                            p = cell.paragraphs[0]
                            if p.runs:
                                p.runs[0].font.name = 'Times New Roman'
                                p.runs[0].font.size = Pt(11)
                            # Alternating zebra background
                            if r_idx % 2 == 1:
                                set_cell_background(cell, "F9FAFC")
                doc.add_paragraph() # Spacer after table

        # Empty line
        if not stripped:
            i += 1
            continue

        # Headings
        if stripped.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(stripped[2:])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x00, 0x33, 0x66) # Deep Navy
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
        elif stripped.startswith("## "):
            p = doc.add_paragraph()
            run = p.add_run(stripped[3:])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x02, 0x4B, 0x86)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
        elif stripped.startswith("### "):
            p = doc.add_paragraph()
            run = p.add_run(stripped[4:])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(13.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
        # Blockquotes
        elif stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(stripped[2:])
            run.font.italic = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
        # Bullet list
        elif stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            clean_text = stripped[2:]
            _render_markdown_inline(p, clean_text)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.25
        # Numbered list
        elif re.match(r'^\d+\.\s', stripped):
            match = re.match(r'^\d+\.\s(.*)', stripped)
            p = doc.add_paragraph(style='List Number')
            _render_markdown_inline(p, match.group(1))
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.25
        # Regular paragraph
        else:
            p = doc.add_paragraph()
            _render_markdown_inline(p, stripped)

        i += 1

    # Save output
    out_path = Path(docx_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"[SUCCESS] Exported: {md_path} -> {docx_path} ({out_path.stat().st_size:,} bytes)")
    return True

def _render_markdown_inline(paragraph, text):
    """Renders bold, italic, and clean links in a paragraph run."""
    # Simple regex parsing for **bold** and *italic*
    tokens = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|\[[^\]]+\]\([^)]+\))', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.font.bold = True
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.font.italic = True
        elif token.startswith("[") and "]" in token and "(" in token and token.endswith(")"):
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', token)
            if m:
                run = paragraph.add_run(m.group(1))
                run.font.underline = True
                run.font.color.rgb = RGBColor(0x02, 0x55, 0x88)
            else:
                paragraph.add_run(token)
        else:
            paragraph.add_run(token)

def export_all_examples():
    """Batch exports all sample deliverables and templates to DOCX."""
    base_dir = Path(__file__).resolve().parent.parent
    examples_dir = base_dir / "examples" / "claims_support_system"
    templates_dir = base_dir / "skills" / "ba-analyst-kit" / "templates"
    dist_dir = base_dir / "dist" / "docx"
    dist_dir.mkdir(parents=True, exist_ok=True)

    targets = list(examples_dir.glob("*.md")) + list(templates_dir.glob("*.md"))
    print(f"=== BATCH EXPORTING {len(targets)} BA DOCUMENTS TO DOCX (NĐ 30/2020) ===")
    
    count = 0
    for t in targets:
        out_name = t.stem + ".docx"
        out_file = dist_dir / out_name
        if convert_md_to_docx(str(t), str(out_file)):
            count += 1
    
    print("-" * 75)
    print(f">>> BATCH EXPORT COMPLETED: {count}/{len(targets)} FILES CONVERTED <<<")
    return count == len(targets)

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Export BA Markdown documents to professional Word DOCX.")
    parser.add_argument("--file", "-f", help="Specific markdown file to convert")
    parser.add_argument("--out", "-o", help="Output docx file path")
    parser.add_argument("--all", "-a", action="store_true", help="Batch export all deliverables and templates")

    args = parser.parse_args()

    if args.all or (not args.file and not args.out):
        success = export_all_examples()
        sys.exit(0 if success else 1)
    elif args.file and args.out:
        success = convert_md_to_docx(args.file, args.out)
        sys.exit(0 if success else 1)
    else:
        parser.print_help()
        sys.exit(1)
